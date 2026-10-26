import csv
import json
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from fastapi.testclient import TestClient

from honeycomb_studio.api import make_app
from honeycomb_studio.data import archive
from honeycomb_studio.engine import Engine
from honeycomb_studio.models import Axis, BatchSpec
from honeycomb_studio.reports import batch_report, incomplete_report, run_report
from honeycomb_studio.store import Store


def wait_operation(client, response):
    assert response.status_code == 200
    identity = response.json()["id"]
    for _ in range(200):
        op = client.get("/v1/operations/" + identity).json()
        if op["status"] != "running":
            assert op["status"] == "completed", op
            return op["result"]
        time.sleep(0.02)
    raise AssertionError("Export did not finish")


def test_batch_reports_preserve_vietnamese_inputs_partial_results_and_black_text(tmp_path):
    store = Store(tmp_path)
    batch = store.create(BatchSpec(name="Quét năng lượng", axes=[Axis(parameter="energy_J", values=[1, 2])]))
    runs = store.list("runs", batch["id"])
    runs[0].update(status="completed", summary={"SEA_J_kg": 12}, quality={"ml_eligible": True})
    assert Path(batch_report(batch, runs)).read_bytes().startswith(b"%PDF")
    file = Path(batch_report(batch, runs, format="docx"))
    document = Document(file)
    text = "\n".join(p.text for p in document.paragraphs)
    text += "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    for expected in (
        "Quét năng lượng",
        "Chiều cao lõi",
        "Mô đun đàn hồi",
        "Thời gian vật lý",
        "Chưa có dữ liệu",
        "Lượt 2",
        "Năng lượng tiêu tán riêng SEA",
    ):
        assert expected in text
    with zipfile.ZipFile(file) as z:
        for name in ("word/document.xml", "word/styles.xml"):
            root = ET.fromstring(z.read(name))
            ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            if name.endswith("document.xml"):
                assert all(e.get(ns + "val") == "000000" for e in root.iter(ns + "color"))
    assert Path(incomplete_report(runs[1], format="docx")).exists()


def test_smoke_report_has_explicit_scope_even_when_opened_without_release_context(tmp_path):
    store = Store(tmp_path)
    batch = store.create(BatchSpec(name="Native environment qualification", data_profile="smoke"))
    runs = store.list("runs", batch["id"])
    file = batch_report(batch, runs, format="docx")
    text = "\n".join(p.text for p in Document(file).paragraphs)
    assert "SMOKE — KIỂM TRA MÔI TRƯỜNG" in text
    assert "không phải dữ liệu nghiên cứu" in text
    assert "không dùng cho huấn luyện" in text
    store.db.close()


def test_batch_dataset_endpoint_exports_all_eligible_runs_and_exclusion_audit(tmp_path):
    store = Store(tmp_path / "data")
    batch = store.create(BatchSpec(axes=[Axis(parameter="energy_J", values=[1, 2, 3, 4])]))
    runs = store.list("runs", batch["id"])
    for run in runs:
        run.update(status="completed", summary={"SEA_J_kg": 12}, quality={"ml_eligible": True})
    runs[2]["quality"]["ml_eligible"] = False
    runs[3]["test_only"] = True
    for run in runs:
        store.save("runs", run)
    app = make_app(store, Engine(store, fake=True), "token", tmp_path / "config", start_worker=False)
    with TestClient(app, headers={"X-Honeycomb-Token": "token"}) as client:
        result = wait_operation(client, client.post(f"/v1/batches/{batch['id']}/dataset/export"))
        assert result["samples"] == 2 and result["excluded_runs"] == 2
        with zipfile.ZipFile(result["path"]) as z:
            samples = list(csv.DictReader(z.read("samples.csv").decode("utf-8-sig").splitlines()))
            coverage = list(csv.DictReader(z.read("job-coverage.csv").decode("utf-8-sig").splitlines()))
            assert {r["run_id"] for r in samples} == {r["id"] for r in runs[:2]}
            assert len(coverage) == 4
            assert {r["split"] for r in samples} == {"train"}
            assert "samples.parquet" in z.namelist() and "data-dictionary.json" in z.namelist()
        response = client.post(f"/v1/batches/{batch['id']}/report/export?format=docx")
        assert Path(wait_operation(client, response)["path"]).suffix == ".docx"
        assert client.post(f"/v1/batches/{batch['id']}/report/export?format=bad").status_code == 422


def test_run_docx_contains_curves_and_editable_tables(tmp_path):
    store = Store(tmp_path)
    batch = store.create(BatchSpec(render_images=False))
    run = store.list("runs", batch["id"])[0]
    run.update(summary={}, quality={"ml_eligible": False})
    attempt = Path(run["path"]) / "attempts/001"
    rows = [
        {
            "time_ms": t,
            "contact_force_N": t * 10,
            "impactor_displacement_mm": t,
            "velocity_mm_s": 100,
            **{
                k + "_J": 1.0
                for k in ("ALLKE", "ALLIE", "ALLPD", "ALLSE", "ALLAE", "ALLVD", "ALLWK", "ETOTAL")
            },
        }
        for t in (0, 1, 2)
    ]
    document = Document(run_report(run, attempt, rows, format="docx"))
    assert len(document.tables) >= 3 and len(document.inline_shapes) >= 3


def test_archive_progress_covers_hashing_and_zip_and_keeps_verified_manifest(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "input.txt").write_text("Mô phỏng", encoding="utf-8")
    events = []
    archive(source, tmp_path / "job.zip", checkpoint=lambda *args: events.append(args))
    assert {e[0] for e in events} == {"hashing", "archiving"}
    assert events[-1][1] == events[-1][2]
    with zipfile.ZipFile(tmp_path / "job.zip") as z:
        assert json.loads(z.read("artifact-manifest.json"))["files"][0]["sha256"]
