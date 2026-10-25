"""Editable Word output from the same content used by the PDF reports."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, Spacer, Table


def save_story(story, destination: Path, landscape=False):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(595.28), Pt(841.89)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Pt(36)
    section.top_margin = section.bottom_margin = Pt(36)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(9 if name == "Normal" else {"Title": 20, "Caption": 8}.get(name, 12))
    available = 841.89 - 72 if landscape else 595.28 - 72

    def emit(items, container, width):
        for item in items:
            if isinstance(item, Paragraph):
                name = item.style.name
                style = {
                    "Heading1": "Heading 1",
                    "Heading2": "Heading 2",
                    "Heading3": "Heading 3",
                    "Title": "Title",
                }.get(name, "Normal")
                paragraph = container.add_paragraph(item.getPlainText(), style)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
            elif isinstance(item, PageBreak):
                container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif isinstance(item, KeepTogether):
                emit(item._content, container, width)
            elif isinstance(item, Image):
                image_width = min(float(item.drawWidth), width)
                container.add_paragraph().add_run().add_picture(str(item.filename), width=Pt(image_width))
            elif isinstance(item, Table):
                grid = container.add_table(rows=len(item._cellvalues), cols=item._ncols)
                grid.style = "Table Grid"
                grid.autofit = False
                source_widths = item._colWidths
                if not all(isinstance(w, (int, float)) for w in source_widths):
                    source_widths = [1] * item._ncols
                total = sum(source_widths)
                for column, size in zip(grid.columns, source_widths, strict=True):
                    column.width = Pt(width * size / total)
                for index, (row, values) in enumerate(zip(grid.rows, item._cellvalues, strict=True)):
                    for cell, value, size in zip(row.cells, values, source_widths, strict=True):
                        cell.width = Pt(width * size / total)
                        children = value if isinstance(value, (list, tuple)) else [value]
                        emit(children, cell, max(20, width * size / total - 12))
                        if not cell.paragraphs[0].text and len(cell.paragraphs) > 1:
                            cell._element.remove(cell.paragraphs[0]._element)
                        if index == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True
                    if index == 0:
                        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                # Prevent Word from splitting a table row across pages.
                for row in grid.rows:
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                # Explicit black text also for table style defaults.
                for row in grid.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.color.rgb = RGBColor(0, 0, 0)
            elif isinstance(item, Spacer):
                container.add_paragraph().paragraph_format.space_after = Pt(min(item.height, 12))
            elif item is not None:
                container.add_paragraph(str(item))

    emit(story, document, available)
    footer = section.footer.paragraphs[0]
    footer.add_run("Honeycomb Abaqus Studio · Trang ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    return str(destination)
